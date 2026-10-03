"""CLI for collecting video metadata, importing VTT files, and searching segments."""
from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
TRANSCRIPTS = DATA / "transcripts"
VIDEO_FIELDS = ["video_id", "title", "published_at", "description", "video_url"]
SEGMENT_FIELDS = ["video_id", "title", "start_sec", "end_sec", "start_time", "text", "hashtags", "video_url", "timestamp_url", "source_file"]
YOUTUBE_API = "https://www.googleapis.com/youtube/v3"


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def youtube_get(endpoint: str, params: dict[str, Any]) -> dict[str, Any]:
    response = requests.get(f"{YOUTUBE_API}/{endpoint}", params=params, timeout=30)
    response.raise_for_status()
    return response.json()


def collect_catalog(channel_id: str) -> None:
    api_key = os.getenv("YOUTUBE_API_KEY")
    if not api_key or api_key == "replace_me":
        raise SystemExit("Set YOUTUBE_API_KEY in .env before collecting the catalogue.")

    channel_data = youtube_get("channels", {"part": "contentDetails,snippet", "id": channel_id, "key": api_key})
    items = channel_data.get("items", [])
    if not items:
        raise SystemExit(f"No channel found for ID: {channel_id}")
    uploads_id = items[0]["contentDetails"]["relatedPlaylists"]["uploads"]
    videos: list[dict[str, str]] = []
    page_token = None

    while True:
        params: dict[str, Any] = {"part": "snippet,contentDetails", "playlistId": uploads_id, "maxResults": 50, "key": api_key}
        if page_token:
            params["pageToken"] = page_token
        page = youtube_get("playlistItems", params)
        for item in page.get("items", []):
            snippet = item.get("snippet", {})
            resource = snippet.get("resourceId", {})
            video_id = resource.get("videoId")
            if not video_id:
                continue
            videos.append({
                "video_id": video_id,
                "title": snippet.get("title", ""),
                "published_at": snippet.get("publishedAt", ""),
                "description": snippet.get("description", ""),
                "video_url": f"https://www.youtube.com/watch?v={video_id}",
            })
        page_token = page.get("nextPageToken")
        if not page_token:
            break

    # Stable de-duplication by video ID.
    unique = {row["video_id"]: row for row in videos}
    rows = sorted(unique.values(), key=lambda row: row["published_at"], reverse=True)
    write_csv(DATA / "videos.csv", rows, VIDEO_FIELDS)
    print(f"Saved {len(rows)} videos to {DATA / 'videos.csv'}")


def parse_vtt_timestamp(value: str) -> float:
    value = value.strip().replace(",", ".")
    parts = value.split(":")
    if len(parts) == 3:
        hours, minutes, seconds = parts
    elif len(parts) == 2:
        hours, (minutes, seconds) = "0", parts
    else:
        raise ValueError(f"Unsupported VTT timestamp: {value}")
    return int(hours) * 3600 + int(minutes) * 60 + float(seconds)


def fmt_time(seconds: float) -> str:
    total = int(seconds)
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def clean_vtt_text(text: str) -> str:
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"&amp;", "&", text)
    text = re.sub(r"&lt;", "<", text)
    text = re.sub(r"&gt;", ">", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_vtt(path: Path, title_by_id: dict[str, str]) -> list[dict[str, str]]:
    video_id = path.stem
    title = title_by_id.get(video_id, video_id)
    content = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = content.splitlines()
    rows: list[dict[str, str]] = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if "-->" not in line:
            i += 1
            continue
        times = line.split("-->", 1)
        try:
            start = parse_vtt_timestamp(times[0].strip().split()[0])
            end = parse_vtt_timestamp(times[1].strip().split()[0])
        except (ValueError, IndexError):
            i += 1
            continue
        i += 1
        text_lines = []
        while i < len(lines) and lines[i].strip():
            candidate = lines[i].strip()
            if not candidate.startswith(("NOTE", "STYLE", "REGION")):
                text_lines.append(candidate)
            i += 1
        text = clean_vtt_text(" ".join(text_lines))
        if not text or end < start:
            continue
        hashtags = " ".join(sorted({tag.lower() for tag in re.findall(r"(?<!\w)#[\w]+", text, flags=re.UNICODE)}))
        url = f"https://www.youtube.com/watch?v={video_id}"
        rows.append({
            "video_id": video_id,
            "title": title,
            "start_sec": f"{start:.3f}",
            "end_sec": f"{end:.3f}",
            "start_time": fmt_time(start),
            "text": text,
            "hashtags": hashtags,
            "video_url": url,
            "timestamp_url": f"{url}&t={int(start)}s",
            "source_file": path.name,
        })
    return rows


def import_transcripts() -> None:
    TRANSCRIPTS.mkdir(parents=True, exist_ok=True)
    videos = read_csv(DATA / "videos.csv")
    title_by_id = {row["video_id"]: row.get("title", row["video_id"]) for row in videos}
    files = sorted(TRANSCRIPTS.glob("*.vtt"))
    if not files:
        raise SystemExit(f"No .vtt files found in {TRANSCRIPTS}. Add authorized subtitle files named VIDEO_ID.vtt.")
    segments: list[dict[str, str]] = []
    for path in files:
        try:
            segments.extend(parse_vtt(path, title_by_id))
            print(f"Imported {path.name}")
        except OSError as exc:
            print(f"Skipped {path.name}: {exc}")
    write_csv(DATA / "segments.csv", segments, SEGMENT_FIELDS)
    print(f"Saved {len(segments)} segments to {DATA / 'segments.csv'}")


def search_segments(query: str) -> None:
    rows = read_csv(DATA / "segments.csv")
    if not rows:
        raise SystemExit("No segments found. Import authorized VTT files first.")
    q = query.strip().casefold()
    exact_phrase = len(q) >= 2 and q.startswith('"') and q.endswith('"')
    if exact_phrase:
        q = q[1:-1].casefold()
    results = []
    for row in rows:
        text = row.get("text", "")
        hashtags = row.get("hashtags", "")
        haystack = (text + " " + hashtags).casefold()
        if q.startswith("#"):
            match = q in hashtags.casefold().split()
        else:
            match = q in haystack
        if match:
            results.append(row)
    write_csv(DATA / "search_results.csv", results, SEGMENT_FIELDS)
    print(f"Found {len(results)} matches. Saved to {DATA / 'search_results.csv'}")
    for row in results[:30]:
        print(f"[{row['start_time']}] {row['title']} — {row['text']}\n  {row['timestamp_url']}\n")



def export_xlsx() -> None:
    """Export the available CSV tables to a formatted Excel workbook."""
    tables = [
        ("Videos", DATA / "videos.csv", VIDEO_FIELDS),
        ("Segments", DATA / "segments.csv", SEGMENT_FIELDS),
        ("SearchResults", DATA / "search_results.csv", SEGMENT_FIELDS),
    ]
    wb = Workbook()
    first = True
    for sheet_name, path, fields in tables:
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        rows = read_csv(path)
        ws.append(fields)
        for row in rows:
            ws.append([row.get(field, "") for field in fields])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col_idx, field in enumerate(fields, start=1):
            values = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
            width = min(max([len(field)] + [len(v) for v in values]) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = max(width, 12)
        if sheet_name in ("Segments", "SearchResults"):
            for row_idx in range(2, ws.max_row + 1):
                link_cell = ws.cell(row=row_idx, column=fields.index("timestamp_url") + 1)
                if link_cell.value:
                    link_cell.hyperlink = link_cell.value
                    link_cell.style = "Hyperlink"
    output = DATA / "youtube_transcripts.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Saved Excel workbook to {output}")

def main() -> None:
    load_dotenv(ROOT / ".env")
    parser = argparse.ArgumentParser(description="Build and search a YouTube transcript catalogue.")
    sub = parser.add_subparsers(dest="command", required=True)
    catalog = sub.add_parser("catalog", help="Collect public metadata for all channel uploads")
    catalog.add_argument("--channel-id", required=True)
    sub.add_parser("import-transcripts", help="Import authorized .vtt files from data/transcripts")
    search = sub.add_parser("search", help="Search imported transcript segments")
    search.add_argument("query", help="Word, phrase, or hashtag; wrap exact phrases in double quotes")
    sub.add_parser("export-xlsx", help="Export videos, transcript segments, and latest search results to Excel")
    args = parser.parse_args()
    if args.command == "catalog":
        collect_catalog(args.channel_id)
    elif args.command == "import-transcripts":
        import_transcripts()
    elif args.command == "search":
        search_segments(args.query)
    elif args.command == "export-xlsx":
        export_xlsx()


if __name__ == "__main__":
    main()
